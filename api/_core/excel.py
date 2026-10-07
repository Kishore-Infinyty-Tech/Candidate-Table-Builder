"""Write the client-ready Excel table."""

from __future__ import annotations

import io
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(name="Calibri", size=11, color="9C0006")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Column widths in Excel character units.
WIDTHS = {
    "Name": 22,
    "No. of Years of Experience": 14,
    "Education (Highest)": 40,
    "Companies (Last 3)": 62,
    "Contact No": 15,
    "Email id": 32,
    "Current CTC (Fixed) + Any variable": 26,
    "Expected CTC (Fixed)": 16,
    "Notice Period/Offer status": 40,
    "Recruiter_Comments": 32,
}
LINE_HEIGHT = 15.0


def _lines_needed(text: str, width: int) -> int:
    usable = max(width - 2, 1)
    return sum(max(1, math.ceil(len(line) / usable)) for line in str(text).split("\n"))


def build_workbook(columns: list[str], rows: list[dict[str, Any]]) -> bytes:
    """rows: [{"values": {column: text}, "highlight": bool}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"

    for ci, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = WIDTHS.get(col, 20)
    ws.row_dimensions[1].height = LINE_HEIGHT * max(
        _lines_needed(c, WIDTHS.get(c, 20)) for c in columns
    ) + 4

    for ri, row in enumerate(rows, 2):
        values = row.get("values") or {}
        red = bool(row.get("highlight"))
        tallest = 1
        for ci, col in enumerate(columns, 1):
            text = str(values.get(col, "") or "")
            cell = ws.cell(row=ri, column=ci, value=text)
            cell.alignment = CENTER
            cell.border = BORDER
            cell.font = RED_FONT if red else BODY_FONT
            cell.number_format = "@"  # keep phone numbers as text
            if red:
                cell.fill = RED_FILL
            tallest = max(tallest, _lines_needed(text, WIDTHS.get(col, 20)))
        ws.row_dimensions[ri].height = LINE_HEIGHT * tallest + 6

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
