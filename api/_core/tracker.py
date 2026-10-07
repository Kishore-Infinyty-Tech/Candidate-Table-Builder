"""Read the tracker workbook and work out which column holds which field.

Tracker column order changes between roles, so columns are found by header
text, never by letter.
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import load_workbook

# Field key -> (label shown in the UI, header matcher).
# Matchers run against a lower-cased, whitespace-collapsed header.
FIELDS: dict[str, tuple[str, Any]] = {
    "name": ("Candidate Name", lambda h: "candidate name" in h or h in ("name", "candidate", "full name")),
    "phone": ("Contact No", lambda h: any(k in h for k in ("contact", "phone", "mobile"))),
    "email": ("Email id", lambda h: "email" in h or "e-mail" in h or "mail id" in h),
    "current_ctc": ("Current CTC (Fixed)", lambda h: "current ctc" in h),
    "variable": ("Variable component", lambda h: "variable" in h),
    "expected_ctc": ("Expected CTC (Fixed)", lambda h: "expected ctc" in h),
    "notice": ("Notice Period", lambda h: "notice" in h),
    "offer": ("Offer status", lambda h: "offer" in h),
    "total_exp": ("Total yrs of experience", lambda h: "total" in h and ("exp" in h or "yrs" in h or "year" in h)),
    "comments": ("Comments", lambda h: "comment" in h or "remark" in h),
}

REQUIRED_FIELDS = ["name", "phone", "email", "current_ctc", "variable", "expected_ctc", "notice", "offer"]


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[ \t]+", " ", str(value)).strip()


def _find_header_row(rows: list[list[Any]]) -> int:
    """The header row is the first of the top rows that names the candidate."""
    for i, row in enumerate(rows[:10]):
        headers = [_norm_header(v) for v in row]
        if any(FIELDS["name"][1](h) for h in headers):
            return i
    return 0


def detect_mapping(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {}
    used: set[int] = set()
    normed = [_norm_header(h) for h in headers]
    for key, (_, matches) in FIELDS.items():
        mapping[key] = None
        for idx, h in enumerate(normed):
            if idx not in used and h and matches(h):
                mapping[key] = idx
                used.add(idx)
                break
    return mapping


def read_tracker(data: bytes) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ValueError("The tracker sheet is empty.")

    header_idx = _find_header_row(rows)
    headers = [_cell_text(h) for h in rows[header_idx]]
    width = len(headers)

    body: list[list[str]] = []
    for raw in rows[header_idx + 1 :]:
        cells = [_cell_text(v) for v in (list(raw) + [None] * width)[:width]]
        if any(cells):
            body.append(cells)

    mapping = detect_mapping(headers)
    # Drop rows with no candidate name (totals, notes, blank formatted rows).
    name_col = mapping.get("name")
    if name_col is not None:
        body = [r for r in body if r[name_col]]

    return {
        "sheet": ws.title,
        "headers": headers,
        "rows": body,
        "mapping": mapping,
        "fields": {k: v[0] for k, v in FIELDS.items()},
        "required": REQUIRED_FIELDS,
    }


def row_value(row: list[str], mapping: dict[str, int | None], field: str) -> str:
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return ""
    return row[idx]
